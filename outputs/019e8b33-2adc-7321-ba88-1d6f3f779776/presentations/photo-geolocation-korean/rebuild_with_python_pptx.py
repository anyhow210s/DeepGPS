from __future__ import annotations

from pathlib import Path
import sys

sys.path.insert(0, str(Path("/Users/koscom/Documents/DeepGPS/.pptx_deps")))

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

from build_pptx import speaker_notes


ROOT = Path("/Users/koscom/Documents/DeepGPS")
OUT = ROOT / "outputs" / "photo_geolocation_presentation_korean.pptx"
TRAINING_FIG = ROOT / "outputs" / "figures" / "training_history.png"
CONFUSION_FIG = ROOT / "outputs" / "figures" / "confusion_matrix.png"

PAPER = RGBColor(247, 243, 234)
INK = RGBColor(23, 33, 43)
MUTED = RGBColor(94, 105, 117)
BLUE = RGBColor(47, 107, 255)
AMBER = RGBColor(244, 163, 64)
GREEN = RGBColor(46, 139, 87)
RED = RGBColor(217, 93, 93)
WHITE = RGBColor(255, 255, 255)
LINE = RGBColor(217, 209, 196)

FONT = "Apple SD Gothic Neo"


def set_fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = color


def add_box(slide, x, y, w, h, fill=WHITE, line=LINE):
    shape = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    return shape


def add_text(slide, x, y, w, h, text, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(0.04)
    frame.margin_right = Inches(0.04)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    frame.vertical_anchor = MSO_ANCHOR.TOP
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_multiline(slide, x, y, w, h, lines, size=14, color=INK, title_color=BLUE):
    box = add_box(slide, x, y, w, h)
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(0.16)
    frame.margin_right = Inches(0.16)
    frame.margin_top = Inches(0.12)
    frame.margin_bottom = Inches(0.12)
    for i, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        run = paragraph.add_run()
        run.text = line
        run.font.name = FONT
        run.font.size = Pt(size + 2 if i == 0 else size)
        run.font.bold = i == 0
        run.font.color.rgb = title_color if i == 0 else color
    return box


def add_bullets(slide, x, y, w, h, title, bullets, accent=BLUE):
    lines = [title] + [f"• {bullet}" for bullet in bullets]
    return add_multiline(slide, x, y, w, h, lines, title_color=accent)


def add_metric(slide, x, y, w, h, value, label, note, accent=BLUE):
    box = add_box(slide, x, y, w, h)
    frame = box.text_frame
    frame.clear()
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, (text, size, color, bold) in enumerate(
        [(value, 25, accent, True), (label, 12, INK, True), (note, 9, MUTED, False)]
    ):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        paragraph.alignment = PP_ALIGN.CENTER
        run = paragraph.add_run()
        run.text = text
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return box


def title(slide, kicker, claim):
    add_text(slide, 0.7, 0.35, 2.5, 0.3, kicker, 11, BLUE, True)
    add_text(slide, 0.7, 0.75, 9.5, 0.75, claim, 28, INK, True)


def footer(slide, n):
    add_text(slide, 0.7, 7.18, 11.9, 0.2, f"Photo Geolocation Prediction using Deep Learning   |   {n}/10", 8, MUTED)


def set_notes(slide, lines):
    text_frame = slide.notes_slide.notes_text_frame
    text_frame.clear()
    text_frame.text = "\n".join(lines)


def blank(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = PAPER
    return slide


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    notes = speaker_notes()

    # 1
    slide = blank(prs)
    slide.background.fill.fore_color.rgb = INK
    add_text(slide, 0.75, 0.7, 7.5, 1.2, "Photo Geolocation\nPrediction", 40, WHITE, True)
    add_text(slide, 0.78, 2.25, 7, 0.4, "using Deep Learning", 24, AMBER, True)
    add_text(slide, 0.78, 2.9, 7.4, 0.7, "EXIF GPS 없이 사진 자체의 시각적 특징으로 위치 카테고리를 예측하는 ResNet50 전이학습 프로젝트", 17, WHITE)
    add_metric(slide, 0.8, 5.0, 2.1, 1.05, "17", "위치 클래스", "수동 분류된 로컬 사진", BLUE)
    add_metric(slide, 3.15, 5.0, 2.1, 1.05, "427", "전체 이미지", "train/val/test 분할", AMBER)
    add_metric(slide, 5.5, 5.0, 2.1, 1.05, "74.07%", "Test Accuracy", "학습에 쓰지 않은 테스트셋", GREEN)
    add_text(slide, 9.25, 5.55, 3.2, 0.5, "Demo: Streamlit UI에서 학습 결과와 예측을 시연", 14, WHITE)
    set_notes(slide, notes[0])

    # 2
    slide = blank(prs)
    title(slide, "PROBLEM", "GPS 메타데이터 없이 사진만 보고 위치를 맞춘다")
    add_bullets(slide, 0.7, 2.2, 3.5, 2.55, "문제 정의", ["입력: 사진 한 장", "출력: 위치 클래스별 확률", "GPS/EXIF 정보는 사용하지 않음"], BLUE)
    add_bullets(slide, 4.55, 2.2, 3.4, 2.55, "왜 어려운가?", ["장소 간 시각적 유사성", "날씨/조명/각도 차이", "개인 사진의 데이터 불균형"], AMBER)
    add_bullets(slide, 8.3, 2.2, 3.9, 2.55, "프로젝트 목표", ["작은 데이터셋으로 완성형 파이프라인 구현", "학습, 평가, 예측, 데모 연결", "설명 가능한 Grad-CAM 포함"], GREEN)
    footer(slide, 2)
    set_notes(slide, notes[1])

    # 3
    slide = blank(prs)
    title(slide, "DATASET", "위치별로 수동 분류한 로컬 사진 폴더를 학습 데이터로 만들었다")
    add_metric(slide, 0.75, 2.0, 1.8, 1.0, "290", "Train", "모델 학습용", BLUE)
    add_metric(slide, 2.8, 2.0, 1.8, 1.0, "56", "Validation", "best model 선택", AMBER)
    add_metric(slide, 4.85, 2.0, 1.8, 1.0, "81", "Test", "최종 평가용", GREEN)
    add_metric(slide, 6.9, 2.0, 1.8, 1.0, "17", "Classes", "위치 카테고리", RED)
    add_bullets(slide, 0.75, 3.6, 5.2, 1.9, "전처리", ["JSON과 동영상은 제외", "손상된 이미지는 skip", "224x224로 변환", "ImageNet mean/std 정규화"], BLUE)
    add_bullets(slide, 6.4, 3.6, 5.2, 1.9, "중요한 관찰", ["Golfland USA는 train 110장으로 가장 많음", "일부 클래스는 test 2~5장뿐", "결과 해석 시 클래스 불균형 고려"], RED)
    footer(slide, 3)
    set_notes(slide, notes[2])

    # 4
    slide = blank(prs)
    title(slide, "PIPELINE", "데이터 준비부터 데모까지 하나의 로컬 딥러닝 시스템으로 연결했다")
    steps = [("Raw\nphotos", "data/raw"), ("Prepare\ndataset", "70/15/15 split"), ("ResNet50", "transfer learning"), ("Evaluate", "accuracy + F1"), ("Demo UI", "Streamlit")]
    for i, (head, sub) in enumerate(steps):
        x = 0.8 + i * 2.4
        add_multiline(slide, x, 2.65, 1.75, 1.15, [head, sub], 12, MUTED, BLUE)
        if i < 4:
            add_text(slide, x + 1.85, 3.0, 0.4, 0.4, "→", 24, BLUE, True, PP_ALIGN.CENTER)
    add_bullets(slide, 1.5, 5.0, 10.2, 0.9, "핵심 원칙", ["온라인 사진 서비스나 외부 AI API를 사용하지 않고, 로컬 폴더와 로컬 PyTorch 학습만 사용한다."], GREEN)
    footer(slide, 4)
    set_notes(slide, notes[3])

    # 5
    slide = blank(prs)
    title(slide, "MODEL", "ImageNet으로 사전학습된 ResNet50을 위치 분류기로 바꿨다")
    layers = [("Input\n224x224", "사진"), ("ResNet50\nBackbone", "시각 특징 추출"), ("New FC\nLayer", "17개 위치 클래스"), ("Softmax", "확률 출력")]
    for i, (head, sub) in enumerate(layers):
        x = 1.0 + i * 2.8
        add_multiline(slide, x, 2.65, 2.15, 1.2, [head, sub], 12, MUTED, [BLUE, BLUE, AMBER, GREEN][i])
        if i < 3:
            add_text(slide, x + 2.25, 3.0, 0.4, 0.4, "→", 24, INK, True, PP_ALIGN.CENTER)
    add_bullets(slide, 1.15, 4.7, 10.9, 1.05, "발표 설명 문장", ["ResNet50은 CNN 기반 이미지 분류 모델이며, 사전학습된 일반 이미지 특징을 활용해 작은 데이터셋에서도 학습을 시작하기 좋습니다."], BLUE)
    footer(slide, 5)
    set_notes(slide, notes[4])

    # 6
    slide = blank(prs)
    title(slide, "TRAINING", "처음에는 분류기만, 이후에는 마지막 ResNet 블록까지 미세조정했다")
    add_multiline(slide, 0.8, 2.05, 5.1, 1.35, ["1단계: Classifier Head", "Epoch 1-5", "Backbone은 고정하고 마지막 분류층만 학습"], 13, MUTED, BLUE)
    add_multiline(slide, 7.0, 2.05, 5.1, 1.35, ["2단계: Fine-tuning", "Epoch 6-8", "ResNet layer4를 열어 위치 데이터에 맞게 조정"], 13, MUTED, AMBER)
    add_bullets(slide, 0.8, 4.2, 5.1, 1.55, "주요 설정", ["Loss: CrossEntropyLoss", "Optimizer: Adam", "Batch size: UI 기본 16", "Best model 기준: validation accuracy"], BLUE)
    add_bullets(slide, 7.0, 4.2, 5.1, 1.55, "왜 이렇게 했나?", ["처음부터 전체 모델을 학습하면 과적합 위험", "분류기 먼저 적응 후 일부 CNN block만 조정", "작은 데이터셋에 적합한 전이학습 전략"], GREEN)
    footer(slide, 6)
    set_notes(slide, notes[5])

    # 7
    slide = blank(prs)
    title(slide, "RESULTS", "검증 정확도는 82.14%, 테스트 정확도는 74.07%로 마무리됐다")
    add_metric(slide, 0.75, 1.85, 1.8, 1.0, "82.14%", "Best Val Acc", "epoch 8 checkpoint", BLUE)
    add_metric(slide, 2.8, 1.85, 1.8, 1.0, "74.07%", "Test Acc", "81 test images", GREEN)
    add_metric(slide, 4.85, 1.85, 1.8, 1.0, "0.68", "Macro F1", "클래스 동일 가중", AMBER)
    add_metric(slide, 6.9, 1.85, 1.8, 1.0, "0.71", "Weighted F1", "샘플 수 반영", RED)
    slide.shapes.add_picture(str(TRAINING_FIG), Inches(1.0), Inches(3.25), width=Inches(5.2))
    add_bullets(slide, 6.7, 3.25, 5.2, 2.2, "해석", ["학습이 진행되며 loss 감소, accuracy 상승", "fine-tuning 이후 validation accuracy 개선", "test accuracy가 validation보다 낮아 일반화 성능은 보수적으로 해석"], GREEN)
    footer(slide, 7)
    set_notes(slide, notes[6])

    # 8
    slide = blank(prs)
    title(slide, "METRICS", "Accuracy만 보지 않고 precision, recall, F1과 혼동 행렬을 함께 봤다")
    slide.shapes.add_picture(str(CONFUSION_FIG), Inches(0.7), Inches(1.9), width=Inches(5.4))
    add_bullets(slide, 6.6, 1.85, 5.4, 0.85, "Accuracy", ["전체 테스트 중 맞춘 비율. 현재 74.07%."], BLUE)
    add_bullets(slide, 6.6, 2.85, 5.4, 0.85, "Precision", ["이 위치라고 예측한 것 중 실제로 맞은 비율."], AMBER)
    add_bullets(slide, 6.6, 3.85, 5.4, 0.85, "Recall", ["실제 해당 위치 사진 중 모델이 찾아낸 비율."], GREEN)
    add_bullets(slide, 6.6, 4.85, 5.4, 0.85, "F1-score", ["Precision과 Recall의 균형. 데이터가 불균형할 때 중요."], RED)
    footer(slide, 8)
    set_notes(slide, notes[7])

    # 9
    slide = blank(prs)
    title(slide, "DEMO", "Streamlit UI에서 학습 실행과 예측 결과 확인까지 시연한다")
    demos = [("1", "Data & Training", "사진 폴더 경로 입력\nPrepare and Train 실행"), ("2", "Training Results", "학습 곡선\nclassification report\nconfusion matrix 확인"), ("3", "Photo Prediction", "새 사진 업로드\nTop-3 예측 확률 확인"), ("4", "Grad-CAM", "모델이 본 영역을\nheatmap으로 설명")]
    for i, (num, head, body) in enumerate(demos):
        x = 0.75 + i * 3.05
        add_multiline(slide, x, 2.5, 2.55, 2.1, [num, head, body], 12, MUTED, BLUE)
    add_multiline(slide, 1.15, 5.45, 11.0, 0.65, ["시연 멘트", "이 UI는 외부 API가 아니라 로컬 모델 파일 best_resnet50.pth를 불러와 예측합니다."], 12, INK, GREEN)
    footer(slide, 9)
    set_notes(slide, notes[8])

    # 10
    slide = blank(prs)
    title(slide, "LESSONS", "작동하는 파이프라인은 만들었지만, 데이터 품질이 성능을 크게 좌우했다")
    add_bullets(slide, 0.75, 2.1, 3.55, 2.45, "잘 된 점", ["ResNet50 전이학습 baseline 구현", "학습/평가/예측/Grad-CAM/UI 연결", "테스트셋 기준 74.07% 달성"], GREEN)
    add_bullets(slide, 4.65, 2.1, 3.55, 2.45, "한계", ["클래스별 이미지 수가 불균형", "일부 클래스 test support가 작음", "비슷한 장소는 혼동 가능"], RED)
    add_bullets(slide, 8.55, 2.1, 3.55, 2.45, "다음 단계", ["클래스별 사진 수 균형 맞추기", "EfficientNetB0 / ViT 비교", "unknown location threshold 추가"], BLUE)
    add_multiline(slide, 1.1, 5.35, 11.1, 0.75, ["결론", "작은 데이터셋에서도 전이학습을 사용하면 완성형 딥러닝 데모를 만들 수 있지만, 실제 성능은 데이터 균형과 테스트셋 크기에 크게 의존한다."], 12, INK, AMBER)
    footer(slide, 10)
    set_notes(slide, notes[9])

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
